"""Compare a manuscript .docx against the generated tables and numbers.

    python src/manuscript_check.py manuscript/private/<draft>.docx [--results results]

Prints every mismatch. Exit status 1 if any. Used by
tests/test_manuscript_numbers.py.

The manuscript is not in the repository (the repo is public and the paper is
unpublished): keep drafts in manuscript/private/ (gitignored) or point
ECOOK_MANUSCRIPT at one. A draft whose wording or table order differs from
manuscript/claims.yaml gets its own manuscript/private/claims.yaml, which is
used in preference, so the patterns stay out of the public repo too.
"""
import argparse
import re
import sys
from pathlib import Path

import pandas as pd
import yaml

ROOT = Path(__file__).resolve().parents[1]


def norm(s):
    return re.sub(r"\s+", " ", str(s).replace(" ", " ")).strip()


def docx_content(path):
    import docx
    d = docx.Document(str(path))
    text = "\n".join(p.text for p in d.paragraphs)
    tables = [[[norm(c.text) for c in r.cells] for r in t.rows] for t in d.tables]
    return text, tables


def claims_path():
    private = ROOT / "manuscript" / "private" / "claims.yaml"
    return private if private.exists() else ROOT / "manuscript" / "claims.yaml"


def compare_tables(doc_tables, gen_tables, order=None):
    """gen_tables: {n: DataFrame}, 1-based. order: the generated table numbers
    in manuscript order (default: all of them, 1, 2, ...). Header rows are
    prose and are not compared; every data cell is. Reported numbers are the
    manuscript's."""
    problems = []
    for pos, n in enumerate(order or sorted(gen_tables), 1):
        gen = gen_tables[n]
        if pos - 1 >= len(doc_tables):
            problems.append(f"Table {pos}: missing from manuscript")
            continue
        doc = doc_tables[pos - 1][1:]
        want = [[norm(v) for v in row] for row in gen.values.tolist()]
        if len(doc) != len(want) or any(len(a) != len(b) for a, b in zip(doc, want)):
            problems.append(f"Table {pos}: shape differs - manuscript "
                            f"{len(doc)}x{len(doc[0]) if doc else 0}, generated "
                            f"{len(want)}x{len(want[0]) if want else 0}. Paste "
                            f"results/tables/table{n}.csv.")
            continue
        for i, (a, b) in enumerate(zip(doc, want)):
            for j, (x, y) in enumerate(zip(a, b)):
                if x != y:
                    problems.append(f"Table {pos} row {i+1} col {j+1}: "
                                    f"manuscript {x!r} != results {y!r}")
    return problems


def compare_claims(text, numbers, claims_file=None):
    """A value key may start with "-" to quote the magnitude of a negative
    result (e.g. a cost change of -5.0% written as "falls by 5.0%")."""
    problems = []
    text = text.replace("\u00a0", " ")
    for c in yaml.safe_load(open(claims_file or claims_path()))["claims"]:
        m = re.search(c["pattern"], text)
        if not m:
            problems.append(f"claim {c['id']}: pattern not found in manuscript "
                            "(text changed? update manuscript/claims.yaml)")
            continue
        for g, (key, fmt) in enumerate(c["values"], 1):
            want = (fmt.format(-numbers[key[1:]]) if key.startswith("-")
                    else fmt.format(numbers[key]))
            if m.group(g) != want:
                problems.append(f"claim {c['id']}: manuscript says {m.group(g)!r}, "
                                f"results give {want!r} ({key})")
    return problems


def check(docx_path, results=ROOT / "results", claims_file=None):
    from manuscript_tables import build_all
    claims_file = claims_file or claims_path()
    order = yaml.safe_load(open(claims_file)).get("tables")
    tables, numbers = build_all(results)
    if order and any(n not in tables for n in order):
        return [f"claims.yaml lists tables {order}; results have {sorted(tables)}"]
    text, doc_tables = docx_content(docx_path)
    return (compare_tables(doc_tables, tables, order)
            + compare_claims(text, numbers, claims_file))


def main():
    sys.path.insert(0, str(ROOT / "src"))
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--results", default=str(ROOT / "results"))
    ap.add_argument("--claims", default=None,
                    help="claims file (default: manuscript/private/claims.yaml "
                         "if present, else manuscript/claims.yaml)")
    a = ap.parse_args()
    probs = check(a.docx, a.results, a.claims)
    print("claims:", a.claims or claims_path())
    for p in probs:
        print("MISMATCH", p)
    print(f"\n{len(probs)} mismatches")
    sys.exit(1 if probs else 0)


if __name__ == "__main__":
    main()
