"""Word count by section for a manuscript .docx (Phase 4: cut to the venue limit).

    python src/word_count.py manuscript/private/<draft>.docx [--limit 5000]

Counts body text only: headings, figure/table captions and the reference list
are reported separately and excluded from the body total, which is how most
conference limits are applied. Check the venue's own rule.
"""
import argparse
import re


def count(path):
    import docx
    d = docx.Document(path)
    sections, cur, refs, captions = {}, "front matter", 0, 0
    in_refs = False
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        n = len(re.findall(r"\S+", t))
        style = (p.style.name if p.style is not None else "").lower()
        if t == "References":
            in_refs = True
            continue
        if in_refs:
            refs += n
            continue
        if "heading" in style or re.match(r"^\d+(\.\d+)?\.? [A-Z]", t) and n <= 12:
            cur = t
            sections.setdefault(cur, 0)
            continue
        if re.match(r"^(Figure|Table) \d+\.", t):
            captions += n
            continue
        sections[cur] = sections.get(cur, 0) + n
    tables = sum(len(re.findall(r"\S+", c.text)) for t in d.tables
                 for r in t.rows for c in r.cells)
    return sections, captions, refs, tables


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--limit", type=int, default=None)
    a = ap.parse_args()
    sections, captions, refs, tables = count(a.docx)
    body = sum(sections.values())
    for k, v in sections.items():
        print(f"  {v:>6}  {k[:70]}")
    print(f"\n  body {body} | captions {captions} | tables {tables} | references {refs}")
    if a.limit:
        print(f"  limit {a.limit}: {'OVER by ' + str(body - a.limit) if body > a.limit else 'within'}")


if __name__ == "__main__":
    main()
