"""Phase 1: every number in every manuscript table matches the results CSVs.

Two parts:
  * test_manuscript_matches_results - the real check. Needs the manuscript
    (manuscript/private/*.docx or $ECOOK_MANUSCRIPT; never committed, the repo
    is public) and real results (results/*.csv from `pipeline.py all`, not
    the synthetic smoke run). Skips, loudly, if either is absent.
  * the comparator itself is tested against a synthetic manuscript, so the
    check is known to work even where it has to skip.
"""
import os
from pathlib import Path

import pytest

from manuscript_check import check, compare_claims, compare_tables, docx_content
from manuscript_tables import build_all

ROOT = Path(__file__).resolve().parents[1]


def _manuscript():
    env = os.environ.get("ECOOK_MANUSCRIPT")
    if env:
        return Path(env)
    hits = sorted((ROOT / "manuscript" / "private").glob("*.docx"))
    return hits[-1] if hits else None


@pytest.mark.manuscript
def test_manuscript_matches_results():
    ms = _manuscript()
    results = Path(os.environ.get("ECOOK_RESULTS", ROOT / "results"))
    if ms is None or not ms.exists():
        pytest.skip("no manuscript: put the draft in manuscript/private/")
    if not (results / "scenarios_both.csv").exists():
        pytest.skip("no real results: run `python src/pipeline.py all`")
    problems = check(ms, results)
    assert not problems, "\n" + "\n".join(problems)


# ------------------------------------------------- the comparator itself
@pytest.fixture(scope="module")
def synthetic_results(tmp_path_factory):
    import subprocess, sys
    out = tmp_path_factory.mktemp("res")
    subprocess.run([sys.executable, str(ROOT / "src" / "pipeline.py"), "phase1",
                    "--synthetic", "--quick", "--results", str(out)],
                   cwd=ROOT, check=True, capture_output=True)
    return out


def _write_docx(path, tables, text):
    import docx
    d = docx.Document()
    d.add_paragraph(text)
    for t in tables.values():
        tab = d.add_table(rows=1, cols=len(t.columns))
        for j, c in enumerate(t.columns):
            tab.rows[0].cells[j].text = str(c)
        for row in t.values.tolist():
            cells = tab.add_row().cells
            for j, v in enumerate(row):
                cells[j].text = str(v)
    d.save(str(path))


def _text(n):
    return (f"Levelised cost falls by {n['lcoe_reduction_pct_rwanda']:.0f}% in Rwanda "
            f"and {n['lcoe_reduction_pct_kenya']:.0f}% in Kenya.")


def test_consistent_manuscript_passes(synthetic_results, tmp_path):
    tables, numbers = build_all(synthetic_results)
    p = tmp_path / "ok.docx"
    _write_docx(p, tables, _text(numbers))
    text, doc_tables = docx_content(p)
    assert compare_tables(doc_tables, tables) == []
    probs = compare_claims(text, numbers)
    # only the claims whose text this stub manuscript lacks may be reported
    assert all("pattern not found" in x for x in probs)
    assert not any("abstract_lcoe_reduction" in x for x in probs)


def test_one_wrong_number_is_caught(synthetic_results, tmp_path):
    tables, numbers = build_all(synthetic_results)
    bad = {k: v.copy() for k, v in tables.items()}
    bad[3].iloc[1, 2] = "$9.99"
    p = tmp_path / "bad.docx"
    _write_docx(p, bad, _text({**numbers, "lcoe_reduction_pct_kenya": 1.0}))
    text, doc_tables = docx_content(p)
    tprobs = compare_tables(doc_tables, tables)
    assert len(tprobs) == 1 and "Table 3 row 2 col 3" in tprobs[0]
    cprobs = compare_claims(text, numbers)
    assert any("abstract_lcoe_reduction" in x and "'1'" in x for x in cprobs)


def test_current_draft_is_out_of_date(synthetic_results):
    """The submitted draft predates Phase 1; it must NOT pass. If this starts
    passing against synthetic results, the check has gone blind."""
    ms = _manuscript()
    if ms is None:
        pytest.skip("no manuscript")
    assert check(ms, synthetic_results)
